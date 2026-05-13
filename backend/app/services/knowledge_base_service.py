import hashlib
import re
import uuid
from pathlib import Path

from fastapi import UploadFile
from sqlmodel import select
from sqlmodel.ext.asyncio.session import AsyncSession

from app.core.logger import logger
from app.models.knowledge_base import (
    KnowledgeBase,
    KnowledgeChunk,
    KnowledgeDocument,
    KnowledgeDocumentStatus,
    KnowledgeIndexJob,
)
from app.schemas.knowledge_base import (
    KnowledgeBaseCreate,
    KnowledgeSearchHit,
    KnowledgeSearchResponse,
)
from app.services.system.file_service import file_service


class KnowledgeBaseService:
    """通用知识库服务。第一版先保证上传、解析、切片和本地检索闭环。"""

    SUPPORTED_EXTENSIONS = {".txt", ".md", ".docx", ".xlsx", ".pdf"}

    async def create(
        self,
        db: AsyncSession,
        obj_in: KnowledgeBaseCreate,
        owner_id: int | None = None,
    ) -> KnowledgeBase:
        collection_name = f"kb_{uuid.uuid4().hex[:16]}"
        db_obj = KnowledgeBase(
            name=obj_in.name,
            description=obj_in.description,
            owner_id=owner_id,
            collection_name=collection_name,
            embedding_model=obj_in.embedding_model,
            is_public=obj_in.is_public,
        )
        db.add(db_obj)
        await db.commit()
        await db.refresh(db_obj)
        return db_obj

    async def list_bases(self, db: AsyncSession, user_id: int | None = None) -> list[KnowledgeBase]:
        statement = select(KnowledgeBase).where(KnowledgeBase.is_deleted == 0)
        result = await db.exec(statement.order_by(KnowledgeBase.create_time.desc()))
        return list(result.all())

    async def upload_document(
        self,
        db: AsyncSession,
        knowledge_base_id: int,
        file: UploadFile,
    ) -> KnowledgeDocument:
        kb = await db.get(KnowledgeBase, knowledge_base_id)
        if not kb or kb.is_deleted:
            raise ValueError("知识库不存在")

        file_name = file.filename or "untitled"
        ext = Path(file_name).suffix.lower()
        if ext not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(f"暂不支持的文件类型: {ext}")

        content = await file.read()
        object_name = f"knowledge_bases/{knowledge_base_id}/{uuid.uuid4().hex}{ext}"
        await file_service.upload_file(
            content,
            object_name,
            content_type=file.content_type or "application/octet-stream",
        )

        doc = KnowledgeDocument(
            knowledge_base_id=knowledge_base_id,
            title=Path(file_name).stem,
            file_name=file_name,
            file_ext=ext.lstrip("."),
            file_path=object_name,
            content_type=file.content_type,
            file_size=len(content),
            status=KnowledgeDocumentStatus.INDEXING,
        )
        db.add(doc)
        await db.commit()
        await db.refresh(doc)

        await self.reindex_document(db, doc.id, file_bytes=content)
        await db.refresh(doc)
        return doc

    async def reindex_document(
        self,
        db: AsyncSession,
        document_id: int,
        file_bytes: bytes | None = None,
    ) -> KnowledgeDocument:
        doc = await db.get(KnowledgeDocument, document_id)
        if not doc or doc.is_deleted:
            raise ValueError("文档不存在")

        job = KnowledgeIndexJob(
            knowledge_base_id=doc.knowledge_base_id,
            document_id=doc.id,
            status=KnowledgeDocumentStatus.INDEXING,
        )
        db.add(job)
        await db.commit()

        try:
            if file_bytes is None:
                file_bytes = await file_service.download_file(doc.file_path)
            text = self._extract_text(doc.file_ext, file_bytes)
            chunks = self._split_text(text)

            old_chunks = await db.exec(
                select(KnowledgeChunk).where(KnowledgeChunk.document_id == doc.id)
            )
            for chunk in old_chunks.all():
                await db.delete(chunk)

            for index, chunk_text in enumerate(chunks):
                vector_id = hashlib.sha1(f"{doc.id}:{index}:{chunk_text}".encode("utf-8")).hexdigest()
                db.add(
                    KnowledgeChunk(
                        knowledge_base_id=doc.knowledge_base_id,
                        document_id=doc.id,
                        chunk_index=index,
                        content=chunk_text,
                        source_label=f"片段 {index + 1}",
                        vector_id=vector_id,
                        token_count=len(chunk_text),
                    )
                )

            doc.status = KnowledgeDocumentStatus.INDEXED
            doc.chunk_count = len(chunks)
            doc.error_message = None
            job.status = KnowledgeDocumentStatus.INDEXED
            job.total_chunks = len(chunks)
            job.indexed_chunks = len(chunks)
            await db.commit()
            logger.info(f"知识库文档已索引: document_id={doc.id}, chunks={len(chunks)}")
            return doc
        except Exception as exc:
            doc.status = KnowledgeDocumentStatus.FAILED
            doc.error_message = str(exc)
            job.status = KnowledgeDocumentStatus.FAILED
            job.error_message = str(exc)
            await db.commit()
            raise

    async def search(
        self,
        db: AsyncSession,
        knowledge_base_id: int,
        query: str,
        top_k: int = 5,
    ) -> KnowledgeSearchResponse:
        statement = select(KnowledgeChunk).where(
            KnowledgeChunk.knowledge_base_id == knowledge_base_id,
            KnowledgeChunk.is_deleted == 0,
        )
        result = await db.exec(statement)
        chunks = list(result.all())
        scored = sorted(
            ((self._score(query, chunk.content), chunk) for chunk in chunks),
            key=lambda item: item[0],
            reverse=True,
        )
        hits: list[KnowledgeSearchHit] = []
        for score, chunk in scored[:top_k]:
            doc = await db.get(KnowledgeDocument, chunk.document_id)
            if not doc:
                continue
            hits.append(
                KnowledgeSearchHit(
                    document_id=doc.id,
                    chunk_id=chunk.id or 0,
                    title=doc.title,
                    content=chunk.content,
                    source_label=chunk.source_label,
                    score=score,
                    metadata=chunk.chunk_metadata,
                )
            )
        return KnowledgeSearchResponse(query=query, hits=hits)

    def _extract_text(self, file_ext: str, file_bytes: bytes) -> str:
        ext = f".{file_ext.lower().lstrip('.')}"
        if ext in {".txt", ".md"}:
            return file_bytes.decode("utf-8", errors="ignore")
        if ext == ".docx":
            import mammoth

            result = mammoth.extract_raw_text(file_bytes)
            return result.value
        if ext == ".xlsx":
            from io import BytesIO

            import openpyxl

            workbook = openpyxl.load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
            lines: list[str] = []
            for sheet in workbook.worksheets:
                lines.append(f"# Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    values = [str(value) for value in row if value is not None]
                    if values:
                        lines.append("\t".join(values))
            return "\n".join(lines)
        if ext == ".pdf":
            return "PDF 文档已上传。第一版保留解析骨架，请接入 pypdf 或企业文档解析服务后重建索引。"
        return ""

    def _split_text(self, text: str, chunk_size: int = 900, overlap: int = 120) -> list[str]:
        normalized = re.sub(r"\n{3,}", "\n\n", text).strip()
        if not normalized:
            return ["文档暂无可解析文本。"]
        chunks: list[str] = []
        start = 0
        while start < len(normalized):
            end = min(start + chunk_size, len(normalized))
            chunks.append(normalized[start:end].strip())
            if end == len(normalized):
                break
            start = max(end - overlap, start + 1)
        return chunks

    def _score(self, query: str, content: str) -> float:
        query_terms = {term for term in re.split(r"\s+", query.lower()) if term}
        if not query_terms:
            return 0.0
        lowered = content.lower()
        hits = sum(1 for term in query_terms if term in lowered)
        return hits / max(len(query_terms), 1)


knowledge_base_service = KnowledgeBaseService()
